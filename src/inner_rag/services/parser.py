"""文档解析：PDF / Word / Excel / 文本 / 图片。

PDF 逐页提取文本，提取不到的页面（扫描件）在启用 OCR 后端时渲染成图片交给 OCR。
"""

from __future__ import annotations

import os
from pathlib import Path

import chardet
from langchain_core.documents import Document
from loguru import logger

from inner_rag.core.config import settings
from inner_rag.services.ocr import OCRBackend, get_ocr_backend

OCR_REQUIRED_HINT = "该文件需要 OCR（扫描件/图片），请先启用 OCR 后端（见 OCR_BACKEND 配置）"


class DocumentParser:
    """通用文档解析器。"""

    SUPPORTED_TYPES: dict[str, list[str]] = {
        "pdf": ["pdf"],
        "word": ["doc", "docx"],
        "excel": ["xls", "xlsx"],
        "text": ["txt", "md", "csv", "json", "xml", "html", "htm"],
        "image": ["jpg", "jpeg", "png", "gif", "bmp", "tiff", "webp"],
    }

    def __init__(self, ocr: OCRBackend | None = None) -> None:
        self._ocr = ocr or get_ocr_backend()

    # ── 公共接口 ───────────────────────────────────────────────────────

    def get_file_type(self, filename: str) -> str:
        ext = Path(filename).suffix.lower().lstrip(".")
        for file_type, extensions in self.SUPPORTED_TYPES.items():
            if ext in extensions:
                return file_type
        return "unknown"

    def parse(self, file_path: str, filename: str | None = None) -> tuple[list[Document], dict]:
        """解析文件，返回 (documents, metadata)。未提取到任何文本时抛 ValueError。"""
        filename = filename or os.path.basename(file_path)
        ext = Path(filename).suffix.lower().lstrip(".")
        file_type = self.get_file_type(filename)
        logger.info(f"[PARSER] 解析 {filename} (type={file_type})")

        if file_type == "pdf":
            documents = self._parse_pdf(file_path)
        elif file_type == "word":
            documents = self._parse_word(file_path)
        elif file_type == "excel":
            documents = self._parse_excel(file_path)
        elif file_type == "text":
            documents = self._parse_text(file_path)
        elif file_type == "image":
            documents = self._parse_image(file_path)
        else:
            msg = f"不支持的文件类型: {ext}"
            raise ValueError(msg)

        documents = [doc for doc in documents if doc.page_content.strip()]
        total_chars = sum(len(doc.page_content) for doc in documents)
        if not documents:
            msg = f"未能从 {filename} 中提取到任何文本"
            if not self._ocr.is_available():
                msg = f"{msg}；{OCR_REQUIRED_HINT}"
            raise ValueError(msg)

        meta = {
            "filename": filename,
            "file_type": file_type,
            "ext": ext,
            "page_count": len(documents),
            "total_chars": total_chars,
            "ocr_enabled": self._ocr.is_available(),
        }
        logger.info(f"[PARSER] {filename}: {len(documents)} 段, {total_chars} 字符")
        return documents, meta

    # ── 各类型实现 ─────────────────────────────────────────────────────

    def _parse_pdf(self, file_path: str) -> list[Document]:
        from pypdf import PdfReader

        reader = PdfReader(file_path)
        documents: list[Document] = []
        mupdf_doc = None
        if self._ocr.is_available():
            import pymupdf

            mupdf_doc = pymupdf.open(file_path)

        try:
            for index, page in enumerate(reader.pages):
                text = (page.extract_text() or "").strip()
                if not text and mupdf_doc is not None:
                    text = self._ocr_pdf_page(mupdf_doc, index)
                if text:
                    documents.append(
                        Document(
                            page_content=text,
                            metadata={"source": file_path, "page": index + 1},
                        )
                    )
        finally:
            if mupdf_doc is not None:
                mupdf_doc.close()

        return documents

    def _ocr_pdf_page(self, mupdf_doc, page_index: int) -> str:
        """把 PDF 页面渲染成图片后交给 OCR 后端。"""
        try:
            page = mupdf_doc[page_index]
            pixmap = page.get_pixmap(dpi=settings.OCR_RENDER_DPI)
            return self._ocr.ocr_bytes(pixmap.tobytes("png"))
        except Exception as exc:
            logger.error(f"[PARSER] PDF 第 {page_index + 1} 页 OCR 失败: {exc}")
            return ""

    def _parse_word(self, file_path: str) -> list[Document]:
        ext = Path(file_path).suffix.lower()
        if ext == ".docx":
            from docx import Document as DocxDocument

            doc = DocxDocument(file_path)
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        paragraphs.append(row_text)
            content = "\n".join(paragraphs)
        else:
            # .doc 是二进制格式，python-docx / docx2txt 都无法直接解析
            try:
                import docx2txt

                content = docx2txt.process(file_path)
            except Exception:
                msg = f"无法解析 .doc 文件（建议先转换为 .docx）: {file_path}"
                raise ValueError(msg) from None

        return [Document(page_content=content, metadata={"source": file_path})]

    def _parse_excel(self, file_path: str) -> list[Document]:
        ext = Path(file_path).suffix.lower()
        documents: list[Document] = []

        if ext == ".xlsx":
            import openpyxl

            workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            try:
                for sheet_name in workbook.sheetnames:
                    rows = []
                    for row in workbook[sheet_name].iter_rows(values_only=True):
                        cells = ["" if cell is None else str(cell) for cell in row]
                        if any(cell.strip() for cell in cells):
                            rows.append(" | ".join(cells))
                    documents.append(
                        Document(
                            page_content=f"[Sheet: {sheet_name}]\n" + "\n".join(rows),
                            metadata={"source": file_path, "sheet": sheet_name},
                        )
                    )
            finally:
                workbook.close()
        else:
            import xlrd

            workbook = xlrd.open_workbook(file_path)
            for sheet in workbook.sheets():
                rows = [
                    " | ".join(str(sheet.cell_value(rx, cx)) for cx in range(sheet.ncols))
                    for rx in range(sheet.nrows)
                ]
                documents.append(
                    Document(
                        page_content=f"[Sheet: {sheet.name}]\n" + "\n".join(rows),
                        metadata={"source": file_path, "sheet": sheet.name},
                    )
                )

        return documents

    def _parse_text(self, file_path: str) -> list[Document]:
        raw = Path(file_path).read_bytes()
        detected = chardet.detect(raw)
        encoding = detected.get("encoding") or "utf-8"
        try:
            content = raw.decode(encoding, errors="replace")
        except (LookupError, UnicodeDecodeError):
            content = raw.decode("utf-8", errors="replace")
        return [Document(page_content=content, metadata={"source": file_path})]

    def _parse_image(self, file_path: str) -> list[Document]:
        if not self._ocr.is_available():
            raise ValueError(f"图片解析需要 OCR：{file_path}；{OCR_REQUIRED_HINT}")

        content = self._ocr.ocr_bytes(Path(file_path).read_bytes())
        return [Document(page_content=content, metadata={"source": file_path, "ocr": True})]


parser = DocumentParser()
